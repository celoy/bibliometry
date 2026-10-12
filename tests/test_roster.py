import openpyxl
import pandas as pd
import pytest

from bibliometry import roster as R


def _workbook(path):
    """Three sheets shaped like the real staff list: header variants, an unlabelled
    e-mail column, ORCID URLs, trailing spaces, one person listed twice."""
    book = openpyxl.Workbook()
    faculty = book.active
    faculty.title = "faculty"
    faculty.append(
        ["Civ", "Nom", "Prénom", "Corps", "Laboratoire", "Unité thématique", None, "ORCID"]
    )
    faculty.append(
        [
            "M.",
            "SOUZA",
            "Ana",
            "PR",
            "WAVES",
            "MECH",
            "a.souza@example.edu",
            "https://orcid.org/0000-0002-1825-0097",
        ]
    )
    faculty.append(
        ["MME", "MARTIN-LEFEVRE", "Helene", "MCF", "OPTICS", "ELEC", "h.ml@example.edu", None]
    )
    teaching = book.create_sheet("Teaching staff")
    teaching.append(["Civ", "Nom", "Prénom", "Laboratoire"])
    teaching.append(["M.", "BERNARD ", "Jean Philippe", "no lab"])
    teaching.append(["M.", "BERNARD ", "Jean Philippe", ""])
    students = book.create_sheet("Doctorants")
    students.append(["Title row that is not a header"])
    students.append(["Nom", "Prénom", "Labo"])
    students.append(["GARCIA LOPEZ", "Pablo", "OLDLAB"])
    book.save(path)


def test_reads_all_sheets(tmp_path):
    path = tmp_path / "list.xlsx"
    _workbook(path)
    roster = R.read_roster(str(path), lab_aliases={"OLDLAB": "NEWLAB"})
    assert list(roster.columns[:9]) == [
        "person_id",
        "last_name",
        "first_name",
        "orcid",
        "openalex_ids",
        "lab",
        "email",
        "sheet",
        "source_row",
    ]
    by_id = roster.set_index("person_id")
    assert by_id.loc["souza--ana", "orcid"] == "0000-0002-1825-0097"
    assert by_id.loc["souza--ana", "email"] == "a.souza@example.edu"  # sniffed, no header
    assert by_id.loc["souza--ana", "lab"] == "WAVES"  # not "Unité thématique"
    assert by_id.loc["garcia-lopez--pablo", "lab"] == "NEWLAB"
    assert by_id.loc["bernard--jean-philippe", "last_name"] == "BERNARD"
    assert roster.attrs["merged"] == ["bernard--jean-philippe"]
    assert len(roster) == 4
    assert "Corps" in roster.columns  # extra columns are kept


def test_sheet_selection(tmp_path):
    path = tmp_path / "list.xlsx"
    _workbook(path)
    assert set(R.read_roster(str(path), sheets=["faculty"]).sheet) == {"faculty"}
    with pytest.raises(R.RosterError, match="Available"):
        R.read_roster(str(path), sheets=["nope"])


def test_csv_with_full_names(tmp_path):
    path = tmp_path / "people.csv"
    path.write_text(
        "Name;Lab;OpenAlex\nSouza, Ana;WAVES;A5000000001\nMARTIN Pierre;MECHLAB;\n"
        "Jeanne Dupont;PROCESS;\n",
        encoding="cp1252",
    )
    roster = R.read_roster(str(path)).set_index("person_id")
    assert roster.loc["souza--ana", "openalex_ids"] == "A5000000001"
    assert "martin--pierre" in roster.index
    assert "dupont--jeanne" in roster.index


def test_column_override(tmp_path):
    path = tmp_path / "people.csv"
    path.write_text("Family;Given\nSOUZA;Ana\n", encoding="utf-8")
    with pytest.raises(R.RosterError, match="no column for names"):
        R.read_roster(str(path))
    roster = R.read_roster(str(path), columns={"last_name": "Family", "first_name": "Given"})
    assert roster.person_id.tolist() == ["souza--ana"]


def test_google_sheet_link_is_exported_as_xlsx(monkeypatch, tmp_path):
    path = tmp_path / "list.xlsx"
    _workbook(path)
    seen = {}

    class Response:
        content = path.read_bytes()

        def raise_for_status(self):
            pass

    def fake_get(url, **_kwargs):
        seen["url"] = url
        return Response()

    monkeypatch.setattr(R.httpx, "get", fake_get)
    roster = R.read_roster("https://docs.google.com/spreadsheets/d/AbC_123/edit#gid=0")
    assert seen["url"].endswith("/d/AbC_123/export?format=xlsx")
    assert len(roster) == 4


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        ("https://orcid.org/0000-0002-1825-009x", "0000-0002-1825-009X"),
        (" 0000-0002-1825-0097 ", "0000-0002-1825-0097"),
        (None, ""),
        ("private", ""),
    ],
)
def test_normalise_orcid(cell, expected):
    assert R.normalise_orcid(cell) == expected


@pytest.mark.parametrize(
    ("full", "expected"),
    [
        ("Souza, Ana", ("Souza", "Ana")),
        ("MARTIN-LEFEVRE Helene", ("MARTIN-LEFEVRE", "Helene")),
        ("Jean-Philippe Bernard", ("Bernard", "Jean-Philippe")),
        ("Madonna", ("Madonna", "")),
    ],
)
def test_split_full_name(full, expected):
    assert R.split_full_name(full) == expected


def test_summary_counts():
    roster = pd.DataFrame(
        {
            "person_id": ["a", "b"],
            "sheet": ["s", "s"],
            "orcid": ["x", ""],
            "openalex_ids": ["", ""],
            "lab": ["L", ""],
        }
    )
    summary = R.summarise_roster(roster)
    assert summary.loc["s"].tolist() == [2, 1, 0, 1]
